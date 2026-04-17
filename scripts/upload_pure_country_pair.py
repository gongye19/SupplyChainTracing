#!/usr/bin/env python3
"""
唯一数据上传脚本（单入口）：
data/pure_country_pair_usd_count/*.xlsx -> country_origin_trade_stats

用法:
  python upload_pure_country_pair.py          # 增量导入
  python upload_pure_country_pair.py --clear  # 先清库再导入
"""

import os
import re
import sys
from datetime import datetime
from pathlib import Path

import openpyxl
import pycountry
from sqlalchemy import create_engine, text

DATABASE_URL = os.getenv(
    "DATABASE_URL",
    os.getenv(
        "DATABASE_PUBLIC_URL",
        "postgresql://postgres:MWXAjkGpQgosJuCgIPcRdudhiyAiXCRl@crossover.proxy.rlwy.net:42314/railway",
    ),
)

BATCH_SIZE = 3000

COUNTRY_NAME_MAP: dict[str, str | None] = {
    "SOUTH KOREA": "KOR",
    "TAIWAN": "TWN",
    "NEW TAIWAN": "TWN",
    "BURMA": "MMR",
    "FORMER REP. OF BIRMAN": "MMR",
    "HOLANDA": "NLD",
    "HRVATSKA": "HRV",
    "SLOVAK REP.": "SVK",
    "DEMOCRATIC PEOPLES REP.": "PRK",
    "U S": "USA",
    "PANAMA": "PAN",
    "KEELING": "CCK",
    "DO SUL": "KOR",
    "SP": "ESP",
    "REPUBLIC": None,
    "REP.": None,
    "IN EUROPE": None,
    "ATTENZA ELECTRONICS": None,
    "ZZZ": None,
}

ISO2_TO_ISO3_MANUAL: dict[str, str] = {
    "UK": "GBR",
    "EL": "GRC",
}

CODE_RE = re.compile(r"\(([A-Z]{2,3})\)")

DDL = """
CREATE TABLE IF NOT EXISTS country_origin_trade_stats (
    hs_code                  VARCHAR(6)  NOT NULL,
    year                     INTEGER     NOT NULL,
    month                    INTEGER     NOT NULL,
    origin_country_code      VARCHAR(3)  NOT NULL,
    destination_country_code VARCHAR(3)  NOT NULL,
    sum_of_usd               NUMERIC(20,2),
    trade_count              INTEGER,
    product_desc_zh          TEXT,
    product_desc_en          TEXT,
    created_at               TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
    updated_at               TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
    CONSTRAINT uq_pair UNIQUE (hs_code, year, month, origin_country_code, destination_country_code)
);
CREATE INDEX IF NOT EXISTS idx_cots_hs ON country_origin_trade_stats(hs_code);
CREATE INDEX IF NOT EXISTS idx_cots_ym ON country_origin_trade_stats(year, month);
CREATE INDEX IF NOT EXISTS idx_cots_origin ON country_origin_trade_stats(origin_country_code);
CREATE INDEX IF NOT EXISTS idx_cots_dest ON country_origin_trade_stats(destination_country_code);
CREATE INDEX IF NOT EXISTS idx_cots_hs_ym ON country_origin_trade_stats(hs_code, year, month);
CREATE INDEX IF NOT EXISTS idx_cots_dest_ym ON country_origin_trade_stats(destination_country_code, year, month);
CREATE INDEX IF NOT EXISTS idx_cots_origin_ym ON country_origin_trade_stats(origin_country_code, year, month);
"""

INSERT_COLS = [
    "hs_code", "year", "month",
    "origin_country_code", "destination_country_code",
    "sum_of_usd", "trade_count", "product_desc_zh", "product_desc_en",
]
KEY_COLS = {"hs_code", "year", "month", "origin_country_code", "destination_country_code"}


def to_iso3(code: str) -> str | None:
    code = (code or "").strip().upper()
    if not code:
        return None
    if len(code) == 3:
        return code
    if len(code) == 2:
        if code in ISO2_TO_ISO3_MANUAL:
            return ISO2_TO_ISO3_MANUAL[code]
        country = pycountry.countries.get(alpha_2=code)
        return country.alpha_3 if country else None
    return None


def clean_country(raw: str) -> str | None:
    if not raw:
        return None
    raw = raw.strip()
    m = CODE_RE.search(raw)
    if m:
        iso3 = to_iso3(m.group(1))
        if iso3:
            return iso3
    up = raw.upper()
    if up in COUNTRY_NAME_MAP:
        return COUNTRY_NAME_MAP[up]
    if re.fullmatch(r"[A-Z]{2,3}", up):
        return to_iso3(up)
    return None


def create_tables(engine) -> None:
    with engine.begin() as conn:
        conn.execute(text(DDL))
    print("✓ 表创建完成")


def clean_database(engine) -> None:
    # 清理当前主表 + 历史表，确保数据库仅保留当前链路所需数据
    drop_tables = [
        "country_origin_trade_stats",
        "country_monthly_trade_stats",
        "shipments_raw",
        "monthly_company_flows",
        "hs_code_categories",
        "port_locations",
        "transactions",
        "companies",
        "locations",
        "categories",
    ]
    with engine.begin() as conn:
        for table in drop_tables:
            conn.execute(text(f'DROP TABLE IF EXISTS "{table}" CASCADE'))
    print("✓ 旧表已清理")


def batch_upsert(engine, batch: list[dict]) -> None:
    if not batch:
        return
    update_cols = [c for c in INSERT_COLS if c not in KEY_COLS]
    set_clause = ", ".join(f"{c} = EXCLUDED.{c}" for c in update_cols)
    all_params: dict = {}
    values_clauses: list[str] = []

    for i, record in enumerate(batch):
        placeholders = ", ".join(f":{c}_{i}" for c in INSERT_COLS)
        values_clauses.append(f"({placeholders})")
        for col in INSERT_COLS:
            all_params[f"{col}_{i}"] = record.get(col)

    cols_str = ", ".join(INSERT_COLS)
    sql = text(
        f"""
        INSERT INTO country_origin_trade_stats ({cols_str})
        VALUES {', '.join(values_clauses)}
        ON CONFLICT (hs_code, year, month, origin_country_code, destination_country_code)
        DO UPDATE SET {set_clause}, updated_at = CURRENT_TIMESTAMP
        """
    )

    with engine.begin() as conn:
        conn.execute(sql, all_params)


def import_xlsx(engine, data_dir: Path) -> int:
    xlsx_files = sorted(data_dir.glob("*.xlsx"))
    if not xlsx_files:
        print(f"❌ 未找到 XLSX 文件: {data_dir}")
        return 0

    print(f"找到 {len(xlsx_files)} 个文件\n")
    total = 0
    skipped = 0

    for fpath in xlsx_files:
        print(f"处理 {fpath.name} ...", end=" ", flush=True)
        t0 = datetime.now()
        wb = openpyxl.load_workbook(fpath, read_only=True)
        ws = wb.active

        batch: list[dict] = []
        count = 0
        file_skipped = 0

        for row in ws.iter_rows(min_row=2, values_only=True):
            hs_code_raw, month_raw, desc_zh, desc_en, imp_raw, exp_raw, usd, cnt = row[:8]
            month_str = str(int(month_raw))
            year = int(month_str[:4])
            month = int(month_str[4:6])

            origin = clean_country(str(exp_raw)) if exp_raw else None
            dest = clean_country(str(imp_raw)) if imp_raw else None
            if not origin or not dest:
                file_skipped += 1
                continue

            batch.append(
                {
                    "hs_code": str(int(hs_code_raw)),
                    "year": year,
                    "month": month,
                    "origin_country_code": origin,
                    "destination_country_code": dest,
                    "sum_of_usd": float(usd) if usd else 0,
                    "trade_count": int(cnt) if cnt else 0,
                    "product_desc_zh": str(desc_zh) if desc_zh else None,
                    "product_desc_en": str(desc_en) if desc_en else None,
                }
            )
            count += 1

            if len(batch) >= BATCH_SIZE:
                batch_upsert(engine, batch)
                batch = []

        if batch:
            batch_upsert(engine, batch)
        wb.close()

        elapsed = (datetime.now() - t0).total_seconds()
        print(f"✓ {count} 条 ({file_skipped} 跳过) [{elapsed:.1f}s]")
        total += count
        skipped += file_skipped

    print(f"\n总计: {total} 条导入, {skipped} 条跳过\n")
    return total


def main() -> None:
    clear_first = "--clear" in sys.argv
    project_root = Path(__file__).resolve().parents[1]
    data_dir = project_root / "data" / "pure_country_pair_usd_count"

    print("=" * 60)
    print("上传 pure_country_pair_usd_count 数据到数据库")
    print("=" * 60)
    print("模式:", "全量重建 (--clear)" if clear_first else "增量导入 (upsert)")
    print()

    if not data_dir.exists():
        print(f"❌ 数据目录不存在: {data_dir}")
        sys.exit(1)

    engine = create_engine(DATABASE_URL, pool_pre_ping=True)
    try:
        with engine.connect() as conn:
            conn.execute(text("SELECT 1"))
        print("✓ 数据库连接成功")
    except Exception as exc:
        print(f"❌ 数据库连接失败: {exc}")
        sys.exit(1)

    if clear_first:
        clean_database(engine)
    create_tables(engine)

    start = datetime.now()
    total = import_xlsx(engine, data_dir)
    elapsed = (datetime.now() - start).total_seconds()

    print("\n" + "=" * 60)
    print("上传完成")
    print("=" * 60)
    print(f"导入记录: {total:,}")
    print(f"耗时: {elapsed:.1f}s")

    with engine.connect() as conn:
        count = conn.execute(text("SELECT COUNT(*) FROM country_origin_trade_stats")).scalar()
    print(f"数据库记录数: {count:,}")


if __name__ == "__main__":
    main()
